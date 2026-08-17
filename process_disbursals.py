import csv
import html
import json
import os
import re
import smtplib
import sys
import tempfile
from collections import defaultdict
from datetime import date, datetime
from email.message import EmailMessage
from pathlib import Path

import boto3
import pymysql

from config import (
    AWS_ACCESS_KEY_ID,
    AWS_REGION,
    AWS_SECRET_ACCESS_KEY,
    DISBURSAL_ALERT_EMAIL_TO,
    S3_BUCKET_LENDER_REPORTS,
    SMTP_FROM,
    SMTP_HOST,
    SMTP_PASSWORD,
    SMTP_PORT,
    SMTP_USER,
    db_config,
)

S3_BUCKET = S3_BUCKET_LENDER_REPORTS
DB_CONFIG = db_config()
EMAIL_FROM = SMTP_FROM
EMAIL_TO = DISBURSAL_ALERT_EMAIL_TO

# Filenames: lenderid_{lender_id}_...
LENDER_CONFIGS = [
    {"lender_id": 1, "status_column": "currentStatus"},
    {"lender_id": 2, "status_column": "Loan Status"},
    {"lender_id": 3, "status_column": "status"},
    {"lender_id": 4, "status_column": "Status"},
    {"lender_id": 7, "status_column": "currentStatus"},
    {"lender_id": 8, "status_column": "Loan Status"},
    {"lender_id": 5, "status_column": "Status"},
    {"lender_id": 6, "status_column": "Status"},
    {"lender_id": 9, "status_column": "Status"},
    {"lender_id": 10, "status_column": "Status"},
    {"lender_id": 11, "status_column": "Status"},
    {"lender_id": 12, "status_column": "Status"},
    {"lender_id": 13, "status_column": "Status"},
    {"lender_id": 14, "status_column": "Status"},
]

# Same lender, different product lines — try both IDs when matching leads/BRE.
LENDER_ID_ALIASES = {
    1: (1, 7),
    7: (7, 1),
    2: (2, 8),
    8: (8, 2),
}

DEFAULT_APP_ID_COLUMNS = (
    "App ID",
    "application_id",
    "Application ID",
    "app_id",
    "ApplicationId",
)
DEFAULT_AMOUNT_COLUMNS = ("Dis Amt", "Disbursed Amount", "dis_amt", "Amount")
DEFAULT_DATE_COLUMNS = ("Dis Date", "Disbursement Date", "dis_date", "Disbursed Date")

# Case-insensitive disbursed statuses.
DISBURSED_STATUSES = frozenset(
    {
        "disbursed",
        "approved process",
    }
)


def file_prefix_for_lender(lender_id):
    return f"lenderid_{int(lender_id)}_"


def lookup_lender_ids(lender_id):
    lender_id = int(lender_id)
    return LENDER_ID_ALIASES.get(lender_id, (lender_id,))


def explain_s3_error(exc):
    message = str(exc)
    if "AWSCompromisedKeyQuarantineV3" in message or "AccessDenied" in message:
        return (
            f"{message}\n\n"
            "This usually means the AWS access key is blocked or lacks S3 permissions.\n"
            "Actions:\n"
            "  1. Create a new IAM access key (the old mf user key may be quarantined by AWS)\n"
            "  2. Ensure the IAM user can s3:ListBucket, s3:GetObject, and s3:DeleteObject on:\n"
            f"     arn:aws:s3:::{S3_BUCKET}\n"
            f"     arn:aws:s3:::{S3_BUCKET}/*\n"
            "  3. Export credentials and rerun:\n"
            "     export AWS_ACCESS_KEY_ID='your-new-key'\n"
            "     export AWS_SECRET_ACCESS_KEY='your-new-secret'\n"
            "     python3 process_disbursals.py"
        )
    return message


def get_s3_client():
    return boto3.client(
        "s3",
        aws_access_key_id=AWS_ACCESS_KEY_ID,
        aws_secret_access_key=AWS_SECRET_ACCESS_KEY,
        region_name=AWS_REGION,
    )


def list_s3_objects(s3_client):
    objects = []
    paginator = s3_client.get_paginator("list_objects_v2")
    for page in paginator.paginate(Bucket=S3_BUCKET):
        for item in page.get("Contents", []):
            objects.append(item["Key"])
    return objects


def match_lender_config(key):
    stem = Path(key).name
    for config in LENDER_CONFIGS:
        if stem.startswith(file_prefix_for_lender(config["lender_id"])):
            return config
    return None


def find_column(fieldnames, column_name):
    if not fieldnames:
        return None
    target = column_name.strip().lower()
    for name in fieldnames:
        if name and name.strip().lower() == target:
            return name
    return None


def find_first_column(fieldnames, column_names):
    for column_name in column_names:
        matched = find_column(fieldnames, column_name)
        if matched:
            return matched
    return None


def app_id_column_candidates(config):
    configured = config.get("app_id_column")
    return (configured,) if configured else DEFAULT_APP_ID_COLUMNS


def normalize_application_id(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "null", "na"}:
        return None
    if text.endswith(".0"):
        text = text[:-2]
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return None


def normalize_status(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "null"}:
        return None
    return text


def normalize_amount(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "null", "na"}:
        return None
    text = text.replace(",", "")
    return text[:15]


def normalize_disbursal_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "null", "na"}:
        return None
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def is_disbursed_status(status):
    return str(status or "").strip().lower() in DISBURSED_STATUSES


def is_empty_criteria_missed(value):
    if value is None:
        return True
    if isinstance(value, (bytes, bytearray)):
        value = value.decode("utf-8", errors="replace")
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return True
        try:
            value = json.loads(text)
        except json.JSONDecodeError:
            text = re.sub(r"\s+", "", text)
            return text in {"[]", "{}", "null", "None"}
    if isinstance(value, dict):
        return len(value) == 0
    if isinstance(value, (list, tuple)):
        return len(value) == 0
    return False


def build_row_record(application_id, status, amount, dis_date):
    return {
        "application_id": application_id,
        "status": status,
        "d_amount": amount,
        "d_date": dis_date,
    }


def read_status_rows_from_csv(file_path, status_column, app_id_columns):
    rows_out = []
    with open(file_path, newline="", encoding="utf-8-sig") as csv_file:
        reader = csv.DictReader(csv_file)
        status_col = find_column(reader.fieldnames, status_column)
        app_id_col = find_first_column(reader.fieldnames, app_id_columns)
        amount_col = find_first_column(reader.fieldnames, DEFAULT_AMOUNT_COLUMNS)
        date_col = find_first_column(reader.fieldnames, DEFAULT_DATE_COLUMNS)

        if not app_id_col:
            print(
                f"  Warning: no application id column "
                f"({', '.join(app_id_columns)}) in {file_path}"
            )
            return rows_out
        if not status_col:
            print(f"  Warning: no '{status_column}' column in {file_path}")

        print(
            f"  Using columns: app_id='{app_id_col}', status='{status_col}', "
            f"amount='{amount_col}', date='{date_col}'"
        )
        for row in reader:
            application_id = normalize_application_id(row.get(app_id_col))
            if application_id is None:
                continue
            status = normalize_status(row.get(status_col) if status_col else None)
            amount = normalize_amount(row.get(amount_col) if amount_col else None)
            dis_date = normalize_disbursal_date(row.get(date_col) if date_col else None)
            rows_out.append(build_row_record(application_id, status, amount, dis_date))
    return rows_out


def read_status_rows_from_xlsx(file_path, status_column, app_id_columns):
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            f"openpyxl is required to read Excel files: {file_path}"
        ) from exc

    rows_out = []
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        workbook.close()
        return rows_out

    header = [str(col).strip() if col is not None else "" for col in header]
    status_col = find_column(header, status_column)
    app_id_col = find_first_column(header, app_id_columns)
    amount_col = find_first_column(header, DEFAULT_AMOUNT_COLUMNS)
    date_col = find_first_column(header, DEFAULT_DATE_COLUMNS)

    if not app_id_col:
        print(
            f"  Warning: no application id column "
            f"({', '.join(app_id_columns)}) in {file_path}"
        )
        workbook.close()
        return rows_out
    if not status_col:
        print(f"  Warning: no '{status_column}' column in {file_path}")

    app_id_index = header.index(app_id_col)
    status_index = header.index(status_col) if status_col else None
    amount_index = header.index(amount_col) if amount_col else None
    date_index = header.index(date_col) if date_col else None
    print(
        f"  Using columns: app_id='{app_id_col}', status='{status_col}', "
        f"amount='{amount_col}', date='{date_col}'"
    )

    for row in rows:
        application_id = normalize_application_id(
            row[app_id_index] if app_id_index < len(row) else None
        )
        if application_id is None:
            continue
        status = normalize_status(
            row[status_index]
            if status_index is not None and status_index < len(row)
            else None
        )
        amount = normalize_amount(
            row[amount_index]
            if amount_index is not None and amount_index < len(row)
            else None
        )
        dis_date = normalize_disbursal_date(
            row[date_index]
            if date_index is not None and date_index < len(row)
            else None
        )
        rows_out.append(build_row_record(application_id, status, amount, dis_date))

    workbook.close()
    return rows_out


def read_status_rows(file_path, config):
    extension = Path(file_path).suffix.lower()
    status_column = config["status_column"]
    app_id_columns = app_id_column_candidates(config)

    if extension == ".csv":
        return read_status_rows_from_csv(file_path, status_column, app_id_columns)
    if extension in {".xlsx", ".xlsm"}:
        return read_status_rows_from_xlsx(file_path, status_column, app_id_columns)

    print(f"  Skipping unsupported file type: {file_path}")
    return []


def fetch_existing_statuses(cursor, lender_id):
    cursor.execute(
        """
        SELECT DISTINCT status
        FROM mf_lender_status_master
        WHERE lender_id = %s
        """,
        (lender_id,),
    )
    return {
        row["status"]
        for row in cursor.fetchall()
        if row["status"] is not None
    }


def insert_new_statuses(cursor, lender_id, statuses):
    first_seen = datetime.now()
    for status in sorted(statuses):
        cursor.execute(
            """
            INSERT INTO mf_lender_status_master (lender_id, status, first_seen)
            VALUES (%s, %s, %s)
            """,
            (lender_id, status, first_seen),
        )


def find_lead_for_update(cursor, application_id, preferred_lender_id):
    for lender_id in lookup_lender_ids(preferred_lender_id):
        cursor.execute(
            """
            SELECT id, lender_id
            FROM lead_master
            WHERE application_id = %s
              AND lender_id = %s
            LIMIT 1
            """,
            (application_id, lender_id),
        )
        row = cursor.fetchone()
        if row:
            return int(row["id"]), int(row["lender_id"])
    return None, None


def is_bre_eligible(cursor, application_id, preferred_lender_id):
    """True if application_bre_logs has empty criteria_missed for preferred/alias lender."""
    for lender_id in lookup_lender_ids(preferred_lender_id):
        cursor.execute(
            """
            SELECT criteria_missed, lender_id
            FROM application_bre_logs
            WHERE application_id = %s
              AND lender_id = %s
            ORDER BY id DESC
            LIMIT 5
            """,
            (application_id, lender_id),
        )
        for row in cursor.fetchall():
            if is_empty_criteria_missed(row.get("criteria_missed")):
                return True, int(row["lender_id"])
    return False, None


def fetch_user_id_for_application(cursor, application_id):
    cursor.execute(
        """
        SELECT userid
        FROM application_master
        WHERE id = %s
        LIMIT 1
        """,
        (application_id,),
    )
    row = cursor.fetchone()
    if not row or row.get("userid") is None:
        return None
    return int(row["userid"])


def update_lead_disburse_status_by_id(cursor, lead_id, disburse_status):
    cursor.execute(
        """
        UPDATE lead_master
        SET disburse_status = %s,
            disbursal_status_check = NOW()
        WHERE id = %s
        """,
        (disburse_status, lead_id),
    )
    return cursor.rowcount


def upsert_mf_disbursal(cursor, user_id, application_id, lender_id, d_status, d_amount, d_date):
    cursor.execute(
        """
        SELECT id
        FROM mf_disbursals
        WHERE application_id = %s
          AND lender_id = %s
        LIMIT 1
        """,
        (application_id, lender_id),
    )
    existing = cursor.fetchone()
    if existing:
        cursor.execute(
            """
            UPDATE mf_disbursals
            SET user_id = COALESCE(%s, user_id),
                d_status = COALESCE(%s, d_status),
                d_amount = COALESCE(%s, d_amount),
                d_date = COALESCE(%s, d_date)
            WHERE id = %s
            """,
            (user_id, d_status, d_amount, d_date, existing["id"]),
        )
        return "updated"
    cursor.execute(
        """
        INSERT INTO mf_disbursals
            (user_id, application_id, lender_id, d_status, d_amount, d_date)
        VALUES (%s, %s, %s, %s, %s, %s)
        """,
        (user_id, application_id, lender_id, d_status, d_amount, d_date),
    )
    return "inserted"


def sync_file_rows(status_rows, preferred_lender_id):
    """Update lead_master and insert mf_disbursals for qualifying rows."""
    updated = 0
    skipped_missing = 0
    bre_accepted = 0
    alias_matched = 0
    disbursals_inserted = 0
    disbursals_updated = 0
    updated_by_lender = defaultdict(int)

    conn = pymysql.connect(**DB_CONFIG)
    try:
        with conn.cursor() as cursor:
            for row in status_rows:
                application_id = row["application_id"]
                status = row["status"]

                lead_id, matched_lender_id = find_lead_for_update(
                    cursor, application_id, preferred_lender_id
                )

                if lead_id is None:
                    bre_ok, bre_lender_id = is_bre_eligible(
                        cursor, application_id, preferred_lender_id
                    )
                    if not bre_ok:
                        skipped_missing += 1
                        continue
                    bre_accepted += 1
                    matched_lender_id = bre_lender_id or int(preferred_lender_id)
                elif matched_lender_id != int(preferred_lender_id):
                    alias_matched += 1

                if lead_id is not None and status:
                    updated += update_lead_disburse_status_by_id(
                        cursor, lead_id, status
                    )
                    updated_by_lender[matched_lender_id] += 1

                # Insert even if Dis Date is empty — nullable fields stay NULL.
                if is_disbursed_status(status):
                    user_id = fetch_user_id_for_application(cursor, application_id)
                    action = upsert_mf_disbursal(
                        cursor,
                        user_id=user_id,
                        application_id=application_id,
                        lender_id=matched_lender_id,
                        d_status=status,
                        d_amount=row.get("d_amount"),
                        d_date=row.get("d_date"),
                    )
                    if action == "inserted":
                        disbursals_inserted += 1
                    else:
                        disbursals_updated += 1

            conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    return {
        "updated": updated,
        "skipped_missing_lead": skipped_missing,
        "bre_accepted": bre_accepted,
        "alias_matched": alias_matched,
        "disbursals_inserted": disbursals_inserted,
        "disbursals_updated": disbursals_updated,
        "updated_by_lender": dict(updated_by_lender),
    }


def sync_statuses_to_db(sheet_statuses, lender_id):
    conn = pymysql.connect(**DB_CONFIG)
    try:
        with conn.cursor() as cursor:
            existing_statuses = fetch_existing_statuses(cursor, lender_id)
            new_statuses = {
                status for status in sheet_statuses if status
            } - existing_statuses
            print(
                f"Existing statuses in DB for lender_id={lender_id}: "
                f"{sorted(existing_statuses)}"
            )
            print(f"New statuses to insert: {sorted(new_statuses)}")
            if new_statuses:
                insert_new_statuses(cursor, lender_id, new_statuses)
                conn.commit()
                print(f"Inserted {len(new_statuses)} new status(es)")
            else:
                print("No new statuses to insert")
            return new_statuses
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def fetch_lender_names(lender_ids=None):
    conn = pymysql.connect(**DB_CONFIG)
    try:
        with conn.cursor() as cursor:
            if lender_ids:
                placeholders = ", ".join(["%s"] * len(lender_ids))
                cursor.execute(
                    f"""
                    SELECT id, lender_name, product_offering
                    FROM mf_lenders
                    WHERE id IN ({placeholders})
                    """,
                    tuple(lender_ids),
                )
            else:
                cursor.execute(
                    """
                    SELECT id, lender_name, product_offering
                    FROM mf_lenders
                    ORDER BY id
                    """
                )
            return {row["id"]: row for row in cursor.fetchall()}
    finally:
        conn.close()


def resolve_lender_name(lender_id, lender_names):
    row = lender_names.get(lender_id)
    if not row:
        return f"Lender {lender_id}"
    name = (row.get("lender_name") or f"Lender {lender_id}").strip()
    offering = (row.get("product_offering") or "").strip()
    if offering:
        return f"{name} - {offering}"
    return name


def fetch_disbursed_counts_by_period():
    """Lender-wise disbursed counts keyed by application_master.created_date."""
    status_list = ", ".join(f"'{status}'" for status in sorted(DISBURSED_STATUSES))
    conn = pymysql.connect(**DB_CONFIG)
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT
                    d.lender_id,
                    COALESCE(NULLIF(TRIM(l.lender_name), ''), CONCAT('Lender ', d.lender_id))
                        AS lender_name,
                    SUM(
                        CASE
                            WHEN a.created_date >= DATE_SUB(
                                CURDATE(), INTERVAL DAYOFWEEK(CURDATE()) - 1 DAY
                            )
                            THEN 1 ELSE 0
                        END
                    ) AS weekly,
                    SUM(
                        CASE
                            WHEN a.created_date >= DATE_FORMAT(CURDATE(), '%%Y-%%m-01')
                            THEN 1 ELSE 0
                        END
                    ) AS monthly,
                    SUM(
                        CASE
                            WHEN a.created_date >= CURDATE() - INTERVAL 30 DAY
                            THEN 1 ELSE 0
                        END
                    ) AS last_30_days,
                    SUM(
                        CASE
                            WHEN a.created_date >= CURDATE() - INTERVAL 90 DAY
                            THEN 1 ELSE 0
                        END
                    ) AS last_3_months,
                    COUNT(*) AS lifetime
                FROM mf_disbursals AS d
                LEFT JOIN application_master AS a ON a.id = d.application_id
                LEFT JOIN mf_lenders AS l ON l.id = d.lender_id
                WHERE LOWER(TRIM(IFNULL(d.d_status, ''))) IN ({status_list})
                GROUP BY d.lender_id, lender_name
                ORDER BY lender_name, d.lender_id
                """
            )
            return cursor.fetchall()
    finally:
        conn.close()


def build_new_status_email(new_statuses_by_lender, lender_names):
    rows = []
    lender_names_for_subject = []
    for lender_id in sorted(new_statuses_by_lender):
        lender_name = resolve_lender_name(lender_id, lender_names)
        lender_names_for_subject.append(lender_name)
        for status in sorted(new_statuses_by_lender[lender_id]):
            rows.append((lender_name, status))

    subject = f"New status code detected - {', '.join(lender_names_for_subject)}"
    html_rows = "".join(
        f"<tr><td>{html.escape(lender_name)}</td><td>{html.escape(status)}</td></tr>"
        for lender_name, status in rows
    )
    html_body = f"""
<html>
  <body>
    <p>The following new lender status codes were detected and inserted:</p>
    <table border="1" cellpadding="8" cellspacing="0">
      <tr><th>Lender Name</th><th>Status</th></tr>
      {html_rows}
    </table>
  </body>
</html>
""".strip()
    text_body = (
        "The following new lender status codes were detected and inserted:\n\n"
        "Lender Name | Status\n"
        + "\n".join(f"{name} | {status}" for name, status in rows)
    )
    return subject, text_body, html_body


def send_email(subject, text_body, html_body, to_emails):
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = EMAIL_FROM
    msg["To"] = ", ".join(to_emails)
    msg.set_content(text_body)
    msg.add_alternative(html_body, subtype="html")
    with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as server:
        server.starttls()
        server.login(SMTP_USER, SMTP_PASSWORD)
        server.send_message(msg)
    print(f"Email sent to {', '.join(to_emails)}: {subject}")


def send_new_status_email(new_statuses_by_lender):
    if not new_statuses_by_lender:
        return
    if not EMAIL_TO:
        print("Skipping new-status email: DISBURSAL_ALERT_EMAIL_TO is empty")
        return
    lender_names = fetch_lender_names(list(new_statuses_by_lender.keys()))
    subject, text_body, html_body = build_new_status_email(
        new_statuses_by_lender, lender_names
    )
    send_email(subject, text_body, html_body, EMAIL_TO)


def build_processing_report_email(summary):
    period_rows = summary["period_disbursals"]

    period_rows_html = "".join(
        (
            "<tr>"
            f"<td>{html.escape(str(row['lender_name']))} (id={row['lender_id']})</td>"
            f"<td>{int(row['weekly'] or 0)}</td>"
            f"<td>{int(row['monthly'] or 0)}</td>"
            f"<td>{int(row['last_30_days'] or 0)}</td>"
            f"<td>{int(row['last_3_months'] or 0)}</td>"
            f"<td>{int(row['lifetime'] or 0)}</td>"
            "</tr>"
        )
        for row in period_rows
    ) or "<tr><td colspan='6'>No disbursed rows in mf_disbursals</td></tr>"

    subject = (
        f"Disbursal processing report - "
        f"{summary['total_api_disbursals']} disbursals - {datetime.now().date()}"
    )

    html_body = f"""
<html>
  <body>
    <h2>Disbursal file processing report</h2>
    <p><strong>Total API Disbursals this run:</strong> {summary['total_api_disbursals']}</p>
    <p><strong>Files processed:</strong> {summary['files_processed']}</p>
    <p><strong>Disbursals Processed:</strong> {summary['rows_read']}</p>
    <p><strong>Skipped (no lead + no BRE eligible):</strong> {summary['skipped_missing_lead']}</p>
    <p><strong>Accepted via BRE (empty criteria_missed):</strong> {summary['bre_accepted']}</p>

    <h3>Actual disbursals in mf_disbursals</h3>
    <table border="1" cellpadding="8" cellspacing="0">
      <tr>
        <th>Lender</th>
        <th>This week</th>
        <th>This month</th>
        <th>Last 30 days</th>
        <th>Last 3 months</th>
        <th>Lifetime</th>
      </tr>
      {period_rows_html}
    </table>
  </body>
</html>
""".strip()

    text_lines = [
        "Disbursal file processing report",
        f"Total API Disbursals this run: {summary['total_api_disbursals']}",
        f"Files processed: {summary['files_processed']}",
        f"Disbursals Processed: {summary['rows_read']}",
        f"Skipped (no lead + no BRE): {summary['skipped_missing_lead']}",
        f"Accepted via BRE: {summary['bre_accepted']}",
        "",
        "Actual disbursals in mf_disbursals:",
    ]
    for row in period_rows:
        text_lines.append(
            f"  {row['lender_name']} (id={row['lender_id']}): "
            f"week={int(row['weekly'] or 0)}, "
            f"month={int(row['monthly'] or 0)}, "
            f"30d={int(row['last_30_days'] or 0)}, "
            f"3m={int(row['last_3_months'] or 0)}, "
            f"lifetime={int(row['lifetime'] or 0)}"
        )
    return subject, "\n".join(text_lines), html_body


def send_processing_report_email(summary):
    if not EMAIL_TO:
        print("Skipping processing report email: DISBURSAL_ALERT_EMAIL_TO is empty")
        return
    subject, text_body, html_body = build_processing_report_email(summary)
    send_email(subject, text_body, html_body, EMAIL_TO)


def local_path_for_key(temp_dir, key):
    return os.path.join(temp_dir, key.replace("/", "_"))


def download_all_files(s3_client, temp_dir):
    keys = list_s3_objects(s3_client)
    downloaded_files = []
    print(f"Found {len(keys)} file(s) in s3://{S3_BUCKET}")
    for key in keys:
        local_path = local_path_for_key(temp_dir, key)
        print(f"Downloading {key}")
        s3_client.download_file(S3_BUCKET, key, local_path)
        downloaded_files.append((key, local_path))
    return downloaded_files


def delete_processed_files_from_s3(s3_client, keys):
    for key in keys:
        print(f"Deleting {key} from s3://{S3_BUCKET}")
        s3_client.delete_object(Bucket=S3_BUCKET, Key=key)
    if keys:
        print(f"Deleted {len(keys)} processed file(s) from S3")


def process_disbursals():
    s3_client = get_s3_client()
    statuses_by_lender = {config["lender_id"]: set() for config in LENDER_CONFIGS}
    rows_by_lender = {config["lender_id"]: [] for config in LENDER_CONFIGS}
    matched_files_by_lender = {config["lender_id"]: 0 for config in LENDER_CONFIGS}
    processed_keys = []

    with tempfile.TemporaryDirectory() as temp_dir:
        downloaded_files = download_all_files(s3_client, temp_dir)

        for key, local_path in downloaded_files:
            config = match_lender_config(key)
            if not config:
                continue

            lender_id = config["lender_id"]
            prefix = file_prefix_for_lender(lender_id)
            matched_files_by_lender[lender_id] += 1
            processed_keys.append(key)
            print(
                f"Processing {key} "
                f"(prefix={prefix}, column={config['status_column']})"
            )
            file_rows = read_status_rows(local_path, config)
            rows_by_lender[lender_id].extend(file_rows)
            file_statuses = {row["status"] for row in file_rows if row.get("status")}
            statuses_by_lender[lender_id].update(file_statuses)
            print(f"  Rows with app_id: {len(file_rows)}")
            print(f"  Unique status values: {sorted(file_statuses)}")

    print()
    processed_any = False
    new_statuses_by_lender = {}
    total_updated = 0
    total_skipped = 0
    total_alias_matched = 0
    total_bre_accepted = 0
    total_disbursals_inserted = 0
    total_disbursals_updated = 0
    rows_read = 0

    for config in LENDER_CONFIGS:
        lender_id = config["lender_id"]
        matched_files = matched_files_by_lender[lender_id]
        all_unique_statuses = statuses_by_lender[lender_id]
        status_rows = rows_by_lender[lender_id]
        prefix = file_prefix_for_lender(lender_id)
        rows_read += len(status_rows)

        if matched_files == 0:
            print(
                f"No files starting with '{prefix}' found in s3://{S3_BUCKET}"
            )
            continue

        processed_any = True
        print(f"Processed {matched_files} {prefix}* file(s)")
        print(
            f"All unique '{config['status_column']}' values from sheet(s): "
            f"{sorted(all_unique_statuses)}"
        )
        print(f"Lookup lender ids (preferred + alias): {lookup_lender_ids(lender_id)}")

        result = sync_file_rows(status_rows, lender_id)
        total_updated += result["updated"]
        total_skipped += result["skipped_missing_lead"]
        total_alias_matched += result["alias_matched"]
        total_bre_accepted += result["bre_accepted"]
        total_disbursals_inserted += result["disbursals_inserted"]
        total_disbursals_updated += result["disbursals_updated"]

        print(
            f"Results for file lender_id={lender_id}: "
            f"lead_updated={result['updated']}, "
            f"skipped={result['skipped_missing_lead']}, "
            f"bre_accepted={result['bre_accepted']}, "
            f"alias_matched={result['alias_matched']}, "
            f"mf_disbursals +{result['disbursals_inserted']}/~{result['disbursals_updated']}"
        )
        print()
        new_statuses = sync_statuses_to_db(all_unique_statuses, lender_id)
        if new_statuses:
            new_statuses_by_lender[lender_id] = new_statuses
        print()

    if not processed_any:
        prefixes = ", ".join(
            file_prefix_for_lender(config["lender_id"]) for config in LENDER_CONFIGS
        )
        print(f"No matching lender files found (expected prefixes: {prefixes})")
        return

    send_new_status_email(new_statuses_by_lender)

    period_disbursals = fetch_disbursed_counts_by_period()
    send_processing_report_email(
        {
            "total_api_disbursals": total_disbursals_inserted + total_disbursals_updated,
            "files_processed": len(processed_keys),
            "rows_read": rows_read,
            "skipped_missing_lead": total_skipped,
            "bre_accepted": total_bre_accepted,
            "period_disbursals": period_disbursals,
        }
    )

    delete_processed_files_from_s3(s3_client, processed_keys)


if __name__ == "__main__":
    try:
        process_disbursals()
    except Exception as exc:
        print(f"Failed to process disbursals: {explain_s3_error(exc)}", file=sys.stderr)
        sys.exit(1)
